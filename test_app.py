"""
Automated Test Suite for DocuMind Local RAG Assistant
Run this script locally to verify your application before pushing to GitHub:
    python test_app.py
"""

import sys
import unittest
from fastapi.testclient import TestClient

from app.main import app
from app.core.rag_engine import (
    extract_youtube_video_id,
    format_timestamp,
    is_simple_greeting,
    GUARDRAIL_REFUSAL_MSG,
    stream_rag_chat
)
from app.config import get_public_settings


class TestDocuMind(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)

    def test_health_check(self):
        """Test API health endpoint"""
        response = self.client.get("/health")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json().get("status"), "healthy")

    def test_get_config(self):
        """Test configuration endpoint returns masked keys and settings"""
        response = self.client.get("/api/config")
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertIn("keys_configured", data)
        self.assertIn("masked_keys", data)
        self.assertIn("strict_guardrails", data)

    def test_youtube_video_id_extraction(self):
        """Test parsing various formats of YouTube URLs"""
        test_cases = [
            ("https://www.youtube.com/watch?v=dQw4w9WgXcQ", "dQw4w9WgXcQ"),
            ("https://youtu.be/dQw4w9WgXcQ", "dQw4w9WgXcQ"),
            ("https://www.youtube.com/embed/dQw4w9WgXcQ", "dQw4w9WgXcQ"),
            ("https://www.youtube.com/shorts/dQw4w9WgXcQ", "dQw4w9WgXcQ"),
            ("https://www.youtube.com/live/dQw4w9WgXcQ", "dQw4w9WgXcQ"),
            ("dQw4w9WgXcQ", "dQw4w9WgXcQ"),
        ]
        for url, expected_id in test_cases:
            with self.subTest(url=url):
                self.assertEqual(extract_youtube_video_id(url), expected_id)

    def test_format_timestamp(self):
        """Test timestamp formatting into [MM:SS]"""
        self.assertEqual(format_timestamp(0), "00:00")
        self.assertEqual(format_timestamp(65), "01:05")
        self.assertEqual(format_timestamp(3665), "01:01:05")

    def test_greeting_guardrail_detection(self):
        """Test simple greetings trigger pleasant greeting responses"""
        self.assertTrue(is_simple_greeting("hello"))
        self.assertTrue(is_simple_greeting("Hi!"))
        self.assertTrue(is_simple_greeting("Good morning"))
        self.assertFalse(is_simple_greeting("What is this video about?"))
        self.assertFalse(is_simple_greeting("Explain quantum computing"))

    def test_strict_guardrail_blocking_empty_vault(self):
        """Test that strict guardrails block off-topic queries when guardrails are enabled"""
        import json
        response = self.client.post(
            "/api/chat/stream",
            json={
                "message": "Who was the first president of the United States?",
                "use_rag": True,
                "strict_guardrails": True
            }
        )
        self.assertEqual(response.status_code, 200)
        accumulated = ""
        for line in response.text.split("\n\n"):
            if line.startswith("data: "):
                try:
                    payload = json.loads(line[6:].strip())
                    if payload.get("type") == "token":
                        accumulated += payload.get("content", "")
                except Exception:
                    pass
        # Must contain guardrail notice or scope refusal
        self.assertTrue("Guardrail Notice" in accumulated or "outside the scope" in accumulated)

    def test_static_index_serving(self):
        """Test root serves HTML application"""
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("DocuMind", response.text)

    def test_scoped_chat_guardrail_missing_scope(self):
        """Test that querying a specific video window scope with no documents returns isolated guardrail notice"""
        import json
        response = self.client.post(
            "/api/chat/stream",
            json={
                "message": "Summarize this video",
                "use_rag": True,
                "strict_guardrails": True,
                "scope": "YouTube: test_missing_video"
            }
        )
        self.assertEqual(response.status_code, 200)
        accumulated = ""
        for line in response.text.split("\n\n"):
            if line.startswith("data: "):
                try:
                    payload = json.loads(line[6:].strip())
                    if payload.get("type") == "token":
                        accumulated += payload.get("content", "")
                except Exception:
                    pass
        self.assertIn("Guardrail Notice", accumulated)
        self.assertIn("test_missing_video", accumulated)

    def test_vector_store_filtered_search(self):
        """Test that Chroma vector search filter isolates chunks by specific document/video"""
        from app.core.vector_store import vector_store_manager
        from langchain_core.documents import Document

        vs = vector_store_manager.get_vector_store()
        doc_a = Document(page_content="Video A talks about Machine Learning and AI agents.", metadata={"filename": "YouTube: video_aaa", "source": "https://youtu.be/video_aaa"})
        doc_b = Document(page_content="Video B talks about French Cooking and Baking Bread.", metadata={"filename": "YouTube: video_bbb", "source": "https://youtu.be/video_bbb"})
        vector_store_manager.add_documents([doc_a, doc_b])

        # Scoped search for Video A only
        res_a = vector_store_manager.search("cooking recipe", k=4, filter={"filename": "YouTube: video_aaa"})
        for doc in res_a:
            self.assertEqual(doc.metadata["filename"], "YouTube: video_aaa")
            self.assertNotIn("French Cooking", doc.page_content)

        # Scoped search for Video B only
        res_b = vector_store_manager.search("machine learning", k=4, filter={"filename": "YouTube: video_bbb"})
        for doc in res_b:
            self.assertEqual(doc.metadata["filename"], "YouTube: video_bbb")
            self.assertNotIn("Machine Learning", doc.page_content)

        # Cleanup test docs
        vector_store_manager.delete_document("YouTube: video_aaa")
        vector_store_manager.delete_document("YouTube: video_bbb")

    def test_delete_endpoint_with_special_characters(self):
        """Test deleting documents with colon/YouTube prefix via DELETE /api/docs/{filename:path}"""
        from app.core.vector_store import vector_store_manager
        from langchain_core.documents import Document

        doc = Document(page_content="Temporary video content", metadata={"filename": "YouTube: test_delete_me"})
        vector_store_manager.add_documents([doc])

        # Test DELETE request with encoded colon & space
        response = self.client.delete("/api/docs/YouTube: test_delete_me")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json().get("success"))

    def test_config_update_resets_vector_cache(self):
        """Test that updating configuration invalidates cached vector store instance"""
        from app.core.vector_store import vector_store_manager
        vector_store_manager.get_vector_store()
        self.assertIsNotNone(vector_store_manager._vector_store)

        # Update a setting via API
        response = self.client.post("/api/config", json={"TEMPERATURE": 0.5})
        self.assertEqual(response.status_code, 200)
        # Vector cache must be cleared for new model/settings to take effect
        self.assertIsNone(vector_store_manager._vector_store)

    def test_pptx_and_xlsx_parsing(self):
        """Test PowerPoint and Excel parsing engines"""
        import os
        import tempfile
        from pptx import Presentation
        from openpyxl import Workbook
        from app.core.rag_engine import load_file_content

        with tempfile.TemporaryDirectory() as tmpdir:
            # 1. Test PPTX
            pptx_path = os.path.join(tmpdir, "test_slides.pptx")
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            slide.shapes.title.text = "DocuMind PPTX Test Slide"
            prs.save(pptx_path)

            pptx_docs = load_file_content(pptx_path)
            self.assertTrue(len(pptx_docs) > 0)
            self.assertIn("DocuMind PPTX Test Slide", pptx_docs[0].page_content)

            # 2. Test XLSX
            xlsx_path = os.path.join(tmpdir, "test_data.xlsx")
            wb = Workbook()
            ws = wb.active
            ws.title = "Financials"
            ws.append(["Category", "Amount"])
            ws.append(["Revenue", 50000])
            wb.save(xlsx_path)

            xlsx_docs = load_file_content(xlsx_path)
            self.assertTrue(len(xlsx_docs) > 0)
            self.assertIn("Financials", xlsx_docs[0].page_content)
            self.assertIn("Revenue", xlsx_docs[0].page_content)

    def test_parse_bool_logic(self):
        """Test robust boolean parsing from strings, booleans, and nulls"""
        from app.config import parse_bool
        self.assertTrue(parse_bool("true"))
        self.assertTrue(parse_bool("1"))
        self.assertTrue(parse_bool(True))
        self.assertFalse(parse_bool("false"))
        self.assertFalse(parse_bool("0"))
        self.assertFalse(parse_bool("off"))
        self.assertFalse(parse_bool(False))
        self.assertTrue(parse_bool(None, default=True))
        self.assertFalse(parse_bool(None, default=False))

    def test_upload_and_delete_documents(self):
        """Test document upload API endpoint and subsequent deletion"""
        import io
        file_content = b"This is a test upload content for DocuMind RAG verification."
        files = [("files", ("unit_test_doc.txt", io.BytesIO(file_content), "text/plain"))]
        
        # Upload
        response = self.client.post("/api/docs/upload", files=files)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(any(u.get("filename") == "unit_test_doc.txt" for u in data.get("uploaded", [])))

        # Delete
        del_response = self.client.delete("/api/docs/unit_test_doc.txt")
        self.assertEqual(del_response.status_code, 200)
        self.assertTrue(del_response.json().get("success"))

    def test_provider_test_connection_invalid_key(self):
        """Test provider key validation handles invalid keys gracefully"""
        response = self.client.post(
            "/api/config/test",
            json={"provider": "groq", "api_key": "gsk_invalid_test_key_12345"}
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertFalse(data.get("success"))

    def test_mask_key_utility(self):
        """Test sensitive API key masking function"""
        from app.config import mask_key
        self.assertEqual(mask_key(""), "")
        self.assertEqual(mask_key("1234"), "••••••••")
        self.assertEqual(mask_key("gsk_1234567890abcdef"), "••••••••cdef")

    def test_stats_with_filename_only_metadata(self):
        """Test that get_stats counts documents whose metadata contains filename without source"""
        from app.core.vector_store import vector_store_manager
        from langchain_core.documents import Document

        doc = Document(page_content="Sample content", metadata={"filename": "filename_only_doc.txt"})
        vector_store_manager.add_documents([doc])
        stats = vector_store_manager.get_stats()
        self.assertTrue(any(d.get("filename") == "filename_only_doc.txt" for d in stats.get("documents", [])))
        # Cleanup
        vector_store_manager.delete_document("filename_only_doc.txt")

    def test_safe_config_parsing(self):
        """Test that get_public_settings never crashes even with null/empty values"""
        from app.config import get_public_settings
        settings = get_public_settings()
        self.assertIsInstance(settings["chunk_size"], int)
        self.assertIsInstance(settings["chunk_overlap"], int)
        self.assertIsInstance(settings["temperature"], float)
        self.assertIsInstance(settings["strict_guardrails"], bool)


if __name__ == "__main__":
    print("=" * 60)
    print("  Running DocuMind Pre-Push Verification Suite...")
    print("=" * 60)
    unittest.main(verbosity=2)
